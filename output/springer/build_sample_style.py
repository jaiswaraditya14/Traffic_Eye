"""Format TrafficEye to the supplied nine-page sample, with blank figure slots."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'D:\Traffic_Eye\output\sample_style')
OUT.mkdir(parents=True, exist_ok=True)
SOURCE = Path(r'D:\Traffic_Eye\output\source\TrafficEye_Sample_Style.md')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.01), Cm(29.706)
sec.left_margin, sec.right_margin = Pt(84), Pt(122)
sec.top_margin, sec.bottom_margin = Pt(97), Pt(151)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def style(name, size=10, bold=False, align=None, before=0, after=6):
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = normal
    s.font.name, s.font.size, s.font.bold = 'Times New Roman', Pt(size), bold
    s.font.color.rgb = RGBColor(0, 0, 0)
    pf = s.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = Pt(size + 2)
    if align is not None:
        pf.alignment = align
    return s

style('Sample Title', 14, True, WD_ALIGN_PARAGRAPH.CENTER, after=18)
style('Sample Authors', 10, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
style('Sample Address', 9, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
style('Sample Abstract', 9, after=12).paragraph_format.left_indent = Pt(20)
style('Sample Keywords', 9, after=18).paragraph_format.left_indent = Pt(20)
style('Sample Section', 12, True, before=18, after=10).paragraph_format.keep_with_next = True
style('Sample Subsection', 10, True, before=12, after=8).paragraph_format.keep_with_next = True
style('Sample Caption', 9, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
style('Sample Reference', 9, after=7).paragraph_format.left_indent = Pt(13)
doc.styles['Sample Reference'].paragraph_format.first_line_indent = Pt(-13)

pending_break = False
section = ''

def paragraph(text, style_name='Normal'):
    global pending_break
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.page_break_before = pending_break
    pending_break = False
    p.add_run(text)
    return p

def blank_figure(height, cols=1):
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = False
    table.rows[0].height = Cm(height)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    trpr = table.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
    for cell in table.rows[0].cells:
        cell.width = Cm(13.74 / cols)
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + edge)
            e.set(qn('w:val'), 'nil')
            borders.append(e)
        tcpr.append(borders)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
    return table

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    i += 1
    if not line:
        continue
    if line == '<!-- PAGE BREAK -->':
        pending_break = True
        continue
    if line.startswith('# '):
        paragraph(line[2:], 'Sample Title')
    elif line.startswith('Meenakshi Garg,'):
        p = paragraph('', 'Sample Authors')
        for n, name in enumerate(line.split(', '), 1):
            if n > 1:
                p.add_run(', ')
            p.add_run(name)
            p.add_run(str(n)).font.superscript = True
    elif line.startswith('Department of'):
        paragraph(line, 'Sample Address')
    elif line.startswith('Correspondence:'):
        p = paragraph('', 'Sample Address')
        for n, email in enumerate(line.split(': ', 1)[1].split('; '), 1):
            if n > 1:
                p.add_run(', ')
            p.add_run(str(n)).font.superscript = True
            p.add_run(email)
        p.paragraph_format.space_after = Pt(22)
    elif line == '## Abstract':
        section = 'Abstract'
    elif line.startswith('### '):
        paragraph(line[4:], 'Sample Subsection')
    elif line.startswith('## '):
        section = line[3:]
        paragraph(section, 'Sample Section')
    elif line.startswith('**Keywords:**'):
        p = paragraph('', 'Sample Keywords')
        p.add_run('Keywords: ').bold = True
        p.add_run(line.split('**Keywords:**', 1)[1].strip().replace(' · ', ', '))
    elif line == '[[FIGURE 1 SPACE]]':
        blank_figure(9.0)
    elif line == '[[FIGURE 2 SPACE]]':
        blank_figure(5.0)
    elif line == '[[FIGURES 3 AND 4 SPACE]]':
        table = blank_figure(8.0, 2)
        captions = []
        while i < len(lines) and len(captions) < 2:
            cap = lines[i].strip()
            i += 1
            if cap.startswith('**Fig.'):
                captions.append(cap.replace('**', ''))
        row = table.add_row()
        for cell, cap in zip(row.cells, captions):
            p = cell.paragraphs[0]
            p.style = doc.styles['Sample Caption']
            label, rest = cap.split('. ', 1)
            p.add_run(label + '. ')
            p.add_run(rest)
    elif line.startswith('**Fig.'):
        paragraph(line.replace('**', ''), 'Sample Caption')
    elif section == 'Abstract':
        p = paragraph('', 'Sample Abstract')
        p.add_run('Abstract. ').bold = True
        p.add_run(line)
    elif section == 'References':
        paragraph(line, 'Sample Reference')
    else:
        p = paragraph(line)
        if re.match(r'^\d+\. ', line):
            label, rest = line.split(': ', 1)
            p.clear()
            p.add_run(label + ': ').bold = True
            p.add_run(rest)

doc.core_properties.title = 'TrafficEye: An Integrated Crowdsourced and AI-Assisted System for Traffic Violation Reporting and Enforcement'
doc.core_properties.author = 'Meenakshi Garg; Shivam Chauhan; Mohit Jain; Aditya Jaiswar'
doc.core_properties.subject = 'Research manuscript formatted against the supplied nine-page sample'
doc.save(OUT / 'TrafficEye_Sample_Style.docx')
print(OUT / 'TrafficEye_Sample_Style.docx')
