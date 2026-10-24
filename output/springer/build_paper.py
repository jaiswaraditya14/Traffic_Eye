from pathlib import Path
import base64
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'D:\Traffic_Eye\output')
SOURCE = ROOT / 'source' / 'TrafficEye_Springer_LNCS.md'
TEMPLATE = ROOT / 'springer' / 'template' / 'splnproc2510.docx'
DEST = ROOT / 'springer' / 'TrafficEye_Springer_LNCS.docx'
ORIGINAL = Path(r'E:\TrafficEye_Rerearch_Paper_i.md')

doc = Document(TEMPLATE)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
odd_header = doc.sections[0].header.paragraphs[0]
even_header = doc.sections[0].even_page_header.paragraphs[0]
for hp, prefix, suffix in (
    (odd_header, '\tTrafficEye: AI-Assisted Traffic Reporting\t', ''),
    (even_header, '', '\tM. Garg et al.'),
):
    hp.clear()
    if prefix:
        hp.add_run(prefix)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    hp._p.append(field)
    if suffix:
        hp.add_run(suffix)

def add(text, style='Normal'):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def diagram(cells):
    table = doc.add_table(rows=1, cols=len(cells))
    for idx, (cell, label) in enumerate(zip(table.rows[0].cells, cells)):
        cell.text = label + (' →' if idx < len(cells) - 1 else '')
        tc_borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            border = OxmlElement('w:' + edge)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            tc_borders.append(border)
        cell._tc.get_or_add_tcPr().append(tc_borders)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

lines = SOURCE.read_text(encoding='utf-8').splitlines()
image_refs = ORIGINAL.read_text(encoding='utf-8').splitlines()
images = {}
for ln in image_refs:
    m = re.match(r'^\[(image[12])\]: <data:image/png;base64,([^>]+)>', ln)
    if m:
        out = ROOT / 'springer' / f'{m.group(1)}.png'
        out.write_bytes(base64.b64decode(m.group(2)))
        images[m.group(1)] = out

section = ''
first_body = False
for i, ln in enumerate(lines):
    line = ln.strip()
    if not line:
        continue
    if line.startswith('# '):
        p = add(line[2:], 'papertitle')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('Meenakshi Garg,'):
        p = add(line, 'author')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith("Vivekanand Education Society") or line.startswith('Correspondence:'):
        p = add(line, 'address')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('## '):
        section = line[3:]
        first_body = True
        if section == 'Abstract':
            continue
        if section in ('Acknowledgments', 'Disclosure of Interests'):
            p = add(section + '.', 'acknowlegments')
            keep_with_next(p)
            continue
        p = add(re.sub(r'^\d+\s+', '', section), 'heading1')
        keep_with_next(p)
        continue
    if line.startswith('**Keywords:**'):
        add(line.replace('**', '').replace('·', ', '), 'keywords')
        continue
    if line.startswith('**Fig. 1.'):
        diagram(['Citizen capture + GPS/EXIF', 'Cloud media + report record', 'Hosted VLM candidate fields', 'Officer verifies'])
        p = add(line.replace('**', ''), 'figurecaption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('**Fig. 2.'):
        diagram(['Submitted media', 'Compress / normalize', 'VLM inference', 'Parse fields', 'Officer review'])
        p = add(line.replace('**', ''), 'figurecaption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('**Fig. 3.'):
        if 'image1' in images:
            p = doc.add_paragraph(style='image')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(images['image1']), height=Cm(4.8))
        p = add(line.replace('**', ''), 'figurecaption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if line.startswith('**Fig. 4.'):
        if 'image2' in images:
            p = doc.add_paragraph(style='image')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(images['image2']), height=Cm(5.2))
        p = add(line.replace('**', ''), 'figurecaption')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if section == 'Abstract':
        add('Abstract. ' + line, 'abstract')
    elif section == 'References':
        add(re.sub(r'^\d+\.\s*', '', line), 'referenceitem')
    elif section in ('Acknowledgments', 'Disclosure of Interests'):
        add(line, 'Normal')
    else:
        add(line, 'p1a' if first_body else 'Normal')
        first_body = False

doc.save(DEST)
print(DEST)
