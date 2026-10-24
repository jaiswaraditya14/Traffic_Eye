from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE

source = Path(r'E:\Research_paper\TrafficEye_Sample_Style.docx')
destination = Path(r'D:\Traffic_Eye\output\sample_style\TrafficEye_Sample_Style_With_Images.docx')
images = [
    Path(r'C:\Users\jaisw\AppData\Local\Temp\codex-clipboard-58b41d99-44a6-4531-8372-270ed90810c0.png'),
    Path(r'C:\Users\jaisw\AppData\Local\Temp\codex-clipboard-5eebd8c8-0400-49e5-b3bf-99a6f874ceb0.png'),
]
captions = [
    'Citizen dashboard showing report totals, verification status, points, and recent reporting activity.',
    'Officer pending-review queue showing advisory severity categories, report details, search, and location filters.',
]
doc = Document(source)
assert len(doc.tables) == 3 and len(doc.tables[2].columns) == 2
assert len(doc.inline_shapes) == 2, 'Preserve the two existing architecture and inference images.'
table = doc.tables[2]
table.rows[0].height = Cm(8.0)
table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
for index, (path, caption) in enumerate(zip(images, captions)):
    cell = table.cell(0, index)
    cell.text = ''
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    shape = p.add_run().add_picture(str(path), height=Cm(7.7))
    shape._inline.docPr.set('name', f'Figure {index + 3}')
    shape._inline.docPr.set('descr', caption)
    cp = table.cell(1, index)
    cp.text = ''
    p = cp.paragraphs[0]
    p.style = doc.styles['Sample Caption']
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f'Fig. {index + 3}. ')
    p.add_run(caption)

old = 'Figures 3 and 4 illustrate the officer queue and an example review screen.'
new = 'Figures 3 and 4 illustrate the citizen dashboard and the officer pending-review queue, respectively.'
found = False
for p in doc.paragraphs:
    if old in p.text:
        p.text = p.text.replace(old, new)
        found = True
assert found, 'The figure discussion must agree with the inserted screenshots.'
doc.save(destination)
assert len(Document(destination).inline_shapes) == 4
with ZipFile(source) as z:
    existing = {sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
with ZipFile(destination) as z:
    embedded = {sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
assert existing.issubset(embedded)
assert all(sha256(p.read_bytes()).hexdigest() in embedded for p in images)
print(destination)
print('Verified: 2 original images retained; 2 screenshots embedded unchanged; Fig. 3/4 labels and discussion corrected.')
